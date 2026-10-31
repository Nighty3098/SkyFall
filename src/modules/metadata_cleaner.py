import os
import shutil
import stat
import tempfile
import zipfile
from pathlib import Path, PurePosixPath, PureWindowsPath

import piexif
from docx import Document
from loguru import logger
from mutagen.flac import FLAC
from mutagen.id3 import ID3, ID3NoHeaderError
from mutagen.mp3 import MP3
from openpyxl import load_workbook
from PIL import Image
from PyPDF2 import PdfReader, PdfWriter

try:
    from pptx import Presentation

    CAN_HANDLE_PPTX = True
except ImportError:
    CAN_HANDLE_PPTX = False

try:
    from odf.opendocument import load as odf_load

    CAN_HANDLE_ODF = True
except ImportError:
    CAN_HANDLE_ODF = False


MAX_ARCHIVE_MEMBERS = 10_000
MAX_ARCHIVE_SIZE = 1_073_741_824
METADATA_SUPPORTED_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".bmp",
    ".tiff",
    ".pdf",
    ".docx",
    ".mp3",
    ".flac",
    ".xlsx",
    ".zip",
    ".pptx",
    ".ppt",
    ".odt",
    ".ods",
    ".epub",
    ".rtf",
}


def _validated_zip_members(zip_ref):
    members = zip_ref.infolist()
    if len(members) > MAX_ARCHIVE_MEMBERS:
        raise ValueError("Archive contains too many entries.")

    total_size = 0
    seen = set()
    validated = []
    for member in members:
        name = member.filename
        path = PurePosixPath(name)
        windows_path = PureWindowsPath(name)
        parts = path.parts
        if (
            not name
            or "\x00" in name
            or ":" in name
            or "\\" in name
            or path.is_absolute()
            or windows_path.is_absolute()
            or windows_path.drive
            or any(part in {"", ".", ".."} for part in name.rstrip("/").split("/"))
        ):
            raise ValueError(f"Unsafe archive member path: {name!r}")

        normalized = "/".join(parts)
        if normalized in seen:
            raise ValueError(f"Duplicate archive member path: {name!r}")
        seen.add(normalized)

        mode = member.external_attr >> 16
        file_type = stat.S_IFMT(mode)
        is_directory = member.is_dir()
        if file_type not in (0, stat.S_IFREG, stat.S_IFDIR):
            raise ValueError(f"Unsupported archive member type: {name!r}")
        if is_directory != (file_type == stat.S_IFDIR) and file_type:
            raise ValueError(f"Invalid archive directory entry: {name!r}")
        if not is_directory:
            total_size += member.file_size
            if total_size > MAX_ARCHIVE_SIZE:
                raise ValueError("Archive exceeds the maximum expanded size.")
        validated.append((member, parts))
    return validated


def _extract_validated_zip(zip_ref, destination, members):
    destination = Path(destination).resolve()
    written_size = 0
    for member, parts in members:
        target = destination.joinpath(*parts)
        if not target.resolve().is_relative_to(destination):
            raise ValueError(
                f"Archive member escapes extraction directory: {member.filename!r}"
            )
        if member.is_dir():
            target.mkdir(parents=True, exist_ok=True)
            continue
        target.parent.mkdir(parents=True, exist_ok=True)
        with zip_ref.open(member) as source, target.open("xb") as output:
            while chunk := source.read(1024 * 1024):
                written_size += len(chunk)
                if written_size > MAX_ARCHIVE_SIZE:
                    raise ValueError("Archive exceeds the maximum expanded size.")
                output.write(chunk)


def _replace_archive_atomically(source_path, build_archive):
    source_path = Path(source_path)
    original_mode = stat.S_IMODE(source_path.stat().st_mode)
    fd, temporary_path = tempfile.mkstemp(
        prefix=f".{source_path.name}.", suffix=".tmp", dir=source_path.parent
    )
    os.close(fd)
    try:
        build_archive(temporary_path)
        os.chmod(temporary_path, original_mode)
        with open(temporary_path, "rb") as temporary_file:
            os.fsync(temporary_file.fileno())
        os.replace(temporary_path, source_path)
    finally:
        if os.path.exists(temporary_path):
            os.unlink(temporary_path)


def _rewrite_archive_from_directory(source_path, members, extracted_root):
    def build_archive(temporary_path):
        with zipfile.ZipFile(temporary_path, "w") as output_zip:
            for member, parts in members:
                output_info = zipfile.ZipInfo(
                    member.filename, date_time=member.date_time
                )
                output_info.compress_type = member.compress_type
                output_info.create_system = member.create_system
                output_info.external_attr = member.external_attr
                if member.is_dir():
                    output_zip.writestr(output_info, b"")
                    continue
                cleaned_path = Path(extracted_root).joinpath(*parts)
                with (
                    cleaned_path.open("rb") as source,
                    output_zip.open(output_info, "w") as output,
                ):
                    shutil.copyfileobj(source, output, length=1024 * 1024)

    _replace_archive_atomically(source_path, build_archive)


def remove_metadata_from_zip(zip_path):
    """
    Unzip -> Remove metadata from each entry -> Re-zip.
    """
    try:
        logger.info(f"Processing ZIP: {zip_path}")
        zip_path = Path(zip_path)
        with tempfile.TemporaryDirectory(prefix="skyfall_zip_") as temp_dir:
            with zipfile.ZipFile(zip_path, "r") as zip_ref:
                members = _validated_zip_members(zip_ref)
                _extract_validated_zip(zip_ref, temp_dir, members)

            for member, parts in members:
                if member.is_dir():
                    continue
                file_path = Path(temp_dir).joinpath(*parts)
                if file_path.suffix.lower() not in METADATA_SUPPORTED_EXTENSIONS:
                    remove_metadata(file_path)
                elif not remove_metadata(file_path):
                    raise RuntimeError(
                        f"Could not clean archive entry: {member.filename}"
                    )

            _rewrite_archive_from_directory(zip_path, members, temp_dir)

        return True
    except Exception as e:
        logger.exception(f"Error processing ZIP {zip_path}: {e}")
        return False


def remove_exif_jpeg(file_path):
    """
    Strip EXIF from JPEG without re-encoding
    """
    try:
        logger.info(f"Stripping EXIF from JPEG: {file_path}")
        piexif.remove(file_path)
        # Verify removal
        exif_dict = piexif.load(file_path)
        if any(exif_dict[tag] for tag in ["0th", "Exif", "GPS", "1st"]):
            logger.info(
                f"EXIF still present after piexif.remove. Re-encoding: {file_path}"
            )
            _reencode_jpeg(file_path)
        else:
            logger.info(f"EXIF successfully removed: {file_path}")
        return True
    except Exception as e:
        logger.exception(f"Error processing JPEG EXIF {file_path}: {e}")
        # Fallback to re-encode if piexif fails
        try:
            _reencode_jpeg(file_path)
            return True
        except Exception as ex:
            logger.exception(f"Failed to re-encode JPEG: {file_path}, error: {ex}")
            return False


def _reencode_jpeg(file_path):
    """
    Re-encode the JPEG with Pillow. This definitely strips EXIF,
    but may alter quality/size if not configured carefully.
    """
    try:
        with Image.open(file_path) as img:
            # Convert to RGB (some images might be in different modes)
            rgb_img = img.convert("RGB")
            # Save, forcing 'exif' to be blank
            rgb_img.save(file_path, format="JPEG", exif=b"")
        logger.info(f"Re-encoded JPEG to remove residual EXIF: {file_path}")
    except Exception as e:
        logger.exception(f"Failed to re-encode JPEG: {file_path}, error: {e}")
        raise


def remove_metadata_from_jpeg(file_path):
    """
    Attempt to remove EXIF from JPEG using piexif (no re-encoding).
    If EXIF still remains, fallback to re-encoding with Pillow to ensure
    all metadata is stripped. This may alter quality/file size slightly.
    """
    try:
        remove_exif_jpeg(file_path)
        return True
    except Exception as e:
        logger.error(f"Failed to remove metadata from JPEG {file_path}: {e}")
        return False


def remove_metadata_from_image(image_path):
    """
    Images beyond JPEG (PNG, GIF, BMP, TIFF):
    Re-encode with Pillow to drop metadata.
    """
    ext = os.path.splitext(image_path)[1].lower()
    if ext in [".jpg", ".jpeg"]:
        return remove_metadata_from_jpeg(image_path)
    try:
        logger.info(f"Processing image: {image_path}")
        with Image.open(image_path) as img:
            data = list(img.getdata())
            clean_img = Image.new(img.mode, img.size)
            clean_img.putdata(data)
            # Keep the original format if possible
            clean_img.save(image_path, format=img.format)
        logger.info(f"Metadata removed from image: {image_path}")
        return True
    except Exception as e:
        logger.exception(f"Error processing image {image_path}: {e}")
        return False


def remove_metadata_from_pdf(pdf_path):
    """
    PDF: Read pages, rewrite them, omit doc info.
    """
    try:
        logger.info(f"Processing PDF: {pdf_path}")
        reader = PdfReader(pdf_path)
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        with open(pdf_path, "wb") as out_pdf:
            writer.write(out_pdf)
        logger.info(f"Metadata removed from PDF: {pdf_path}")
        return True
    except Exception as e:
        logger.exception(f"Error processing PDF {pdf_path}: {e}")
        return False


def remove_metadata_from_docx(docx_path):
    """
    DOCX: Clear core properties with python-docx.
    """
    try:
        logger.info(f"Processing DOCX: {docx_path}")
        doc = Document(docx_path)
        metadata_fields = [
            "author",
            "comments",
            "category",
            "content_status",
            "identifier",
            "keywords",
            "language",
            "last_modified_by",
            "last_printed",
            "revision",
            "subject",
            "title",
            "version",
        ]
        for prop in metadata_fields:
            try:
                setattr(doc.core_properties, prop, "")
            except ValueError:
                pass
        doc.settings.odd_and_even_pages_header_footer = False
        doc.save(docx_path)
        logger.info(f"Metadata removed from DOCX: {docx_path}")
        return True
    except Exception as e:
        logger.exception(f"Error processing DOCX {docx_path}: {e}")
        return False


def remove_metadata_from_pptx(pptx_path):
    """
    PPTX: Clear core properties with python-pptx.
    """
    if not CAN_HANDLE_PPTX:
        logger.warning("python-pptx not installed; cannot process PPTX.")
        return False
    try:
        logger.info(f"Processing PPTX: {pptx_path}")
        ppt = Presentation(pptx_path)
        props = ppt.core_properties
        props.author = ""
        props.category = ""
        props.comments = ""
        props.content_status = ""
        props.created = None
        props.identifier = ""
        props.keywords = ""
        props.last_modified_by = ""
        props.last_printed = None
        props.modified = None
        props.revision = ""
        props.subject = ""
        props.title = ""
        ppt.save(pptx_path)
        logger.info(f"Metadata removed from PPTX: {pptx_path}")
        return True
    except Exception as e:
        logger.exception(f"Error processing PPTX {pptx_path}: {e}")
        return False


def remove_metadata_from_ppt(ppt_path):
    """
    PPT (legacy):
    Typically convert .ppt -> .pptx, then do the PPTX route.
    """
    try:
        logger.warning(
            f".ppt is legacy binary; recommended to convert {ppt_path} to .pptx first."
        )
        return False
    except Exception as e:
        logger.exception(f"Error processing PPT {ppt_path}: {e}")
        return False


def remove_metadata_from_odt(odt_path):
    """
    ODT (OpenDocument Text): Use odfpy to remove <office:meta>.
    """
    if not CAN_HANDLE_ODF:
        logger.warning("odfpy not installed; cannot process ODT.")
        return False
    try:
        logger.info(f"Processing ODT: {odt_path}")
        doc = odf_load(odt_path)
        meta = doc.meta
        # ~(^-^)~ Remove all metadata children
        for child in list(meta.childNodes):
            meta.removeChild(child)
        doc.save(odt_path)
        logger.info(f"Metadata removed from ODT: {odt_path}")
        return True
    except Exception as e:
        logger.exception(f"Error processing ODT {odt_path}: {e}")
        return False


def remove_metadata_from_ods(ods_path):
    """
    ODS (OpenDocument Spreadsheet): Also via odfpy.
    """
    if not CAN_HANDLE_ODF:
        logger.warning("odfpy not installed; cannot process ODS.")
        return False
    try:
        logger.info(f"Processing ODS: {ods_path}")
        spreadsheet = odf_load(ods_path)
        meta = spreadsheet.meta
        for child in list(meta.childNodes):
            meta.removeChild(child)
        spreadsheet.save(ods_path)
        logger.info(f"Metadata removed from ODS: {ods_path}")
        return True
    except Exception as e:
        logger.exception(f"Error processing ODS {ods_path}: {e}")
        return False


def remove_metadata_from_epub(epub_path):
    """
    EPUB: Search for .opf files, remove <metadata> content.
    """
    try:
        logger.info(f"Processing EPUB: {epub_path}")
        epub_path = Path(epub_path)
        with tempfile.TemporaryDirectory(prefix="skyfall_epub_") as temp_dir:
            with zipfile.ZipFile(epub_path, "r") as zip_ref:
                members = _validated_zip_members(zip_ref)
                _extract_validated_zip(zip_ref, temp_dir, members)

            opf_files = [
                Path(temp_dir).joinpath(*parts)
                for member, parts in members
                if not member.is_dir() and member.filename.lower().endswith(".opf")
            ]

            import re

            for opf_file in opf_files:
                with open(opf_file, encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                content = re.sub(
                    r"<metadata[^>]*>.*?</metadata>",
                    "<metadata></metadata>",
                    content,
                    flags=re.DOTALL,
                )
                with open(opf_file, "w", encoding="utf-8") as f:
                    f.write(content)
                logger.info(f"Metadata stripped from OPF file: {opf_file}")
            _rewrite_archive_from_directory(epub_path, members, temp_dir)
        logger.info(f"Re-zipped EPUB after metadata removal: {epub_path}")

        return True
    except Exception as e:
        logger.exception(f"Error processing EPUB {epub_path}: {e}")
        return False


def remove_metadata_from_rtf(rtf_path):
    r"""
    RTF: Naive \info removal with regex.
    """
    try:
        logger.info(f"Processing RTF: {rtf_path}")
        with open(rtf_path, encoding="utf-8", errors="ignore") as f:
            content = f.read()
        import re

        content = re.sub(r"{\\info[^}]*}", "", content, flags=re.DOTALL)
        with open(rtf_path, "w", encoding="utf-8") as f:
            f.write(content)
        logger.info(f"Metadata removed from RTF: {rtf_path}")
        return True
    except Exception as e:
        logger.exception(f"Error processing RTF {rtf_path}: {e}")
        return False


def remove_metadata_from_mp3(mp3_path):
    """
    MP3: Delete ID3 tags with mutagen.
    """
    try:
        logger.info(f"Processing MP3: {mp3_path}")
        audio = MP3(mp3_path, ID3=ID3)
        audio.delete()
        audio.save(mp3_path)
        logger.info(f"Metadata removed from MP3: {mp3_path}")
        return True
    except ID3NoHeaderError:
        try:
            audio = MP3(mp3_path)
            audio.save(mp3_path)
            logger.info(f"No ID3 header found, but file saved: {mp3_path}")
            return True
        except Exception as e:
            logger.exception(f"Error saving MP3 {mp3_path} without ID3 header: {e}")
            return False
    except Exception as e:
        logger.exception(f"Error processing MP3 {mp3_path}: {e}")
        return False


def remove_metadata_from_flac(flac_path):
    """
    FLAC: Remove tags with mutagen.
    """
    try:
        logger.info(f"Processing FLAC: {flac_path}")
        audio = FLAC(flac_path)
        audio.delete()
        audio.save()
        logger.info(f"Metadata removed from FLAC: {flac_path}")
        return True
    except Exception as e:
        logger.exception(f"Error processing FLAC {flac_path}: {e}")
        return False


def remove_metadata_from_xlsx(xlsx_path):
    """
    XLSX: Clear workbook properties with openpyxl.
    """
    try:
        logger.info(f"Processing XLSX: {xlsx_path}")
        workbook = load_workbook(filename=xlsx_path)
        metadata_fields = [
            "creator",
            "title",
            "subject",
            "description",
            "keywords",
            "category",
            "comments",
            "last_modified_by",
            "company",
            "manager",
        ]
        for prop in metadata_fields:
            try:
                setattr(workbook.properties, prop, "")
            except ValueError:
                pass
        workbook.save(xlsx_path)
        logger.info(f"Metadata removed from XLSX: {xlsx_path}")
        return True
    except Exception as e:
        logger.exception(f"Error processing XLSX {xlsx_path}: {e}")
        return False


def remove_metadata(file_path):
    """
    Decide how to remove metadata based on file extension.
    """
    logger.debug(f"File: {file_path}")
    if not isinstance(file_path, (str, bytes, os.PathLike)):
        logger.error(f"Invalid file path type: {type(file_path).__name__}")
        return False
    file_extension = os.path.splitext(file_path)[1].lower()

    if file_extension in [".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff"]:
        return remove_metadata_from_image(file_path)

    elif file_extension == ".pdf":
        return remove_metadata_from_pdf(file_path)

    elif file_extension == ".docx":
        return remove_metadata_from_docx(file_path)

    elif file_extension == ".mp3":
        return remove_metadata_from_mp3(file_path)
    elif file_extension == ".flac":
        return remove_metadata_from_flac(file_path)

    elif file_extension == ".xlsx":
        return remove_metadata_from_xlsx(file_path)

    elif file_extension == ".zip":
        return remove_metadata_from_zip(file_path)

    elif file_extension == ".pptx":
        return remove_metadata_from_pptx(file_path)
    elif file_extension == ".ppt":
        return remove_metadata_from_ppt(file_path)

    elif file_extension == ".odt":
        return remove_metadata_from_odt(file_path)
    elif file_extension == ".ods":
        return remove_metadata_from_ods(file_path)

    elif file_extension == ".epub":
        return remove_metadata_from_epub(file_path)

    elif file_extension == ".rtf":
        return remove_metadata_from_rtf(file_path)

    else:
        logger.warning(f"Unsupported file type {file_extension} for {file_path}.")
        return False
